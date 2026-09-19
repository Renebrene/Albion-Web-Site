from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_DIR = Path(__file__).resolve().parents[2]
INPUT_DIR = PROJECT_DIR / "input"
TMP_DIR = PROJECT_DIR / "tmp"

MONTHS = ["Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]

KNOWN_HEADERS = {
    "Город",
    "Город 1",
    "город 2",
    "Наименование",
    "Адрес",
    "Адресс 2",
    "Сторона",
    "Формат конструкции",
    "Размер",
    "Цена за месяц",
    "Координаты",
    "Свет",
    "Материал",
    "№ конструкции",
    "Ссылка на фото",
    "Карта",
    *MONTHS,
}


def clean(value: Any) -> str:
    if value is None:
        return ""

    return str(value).strip()


def normalize_header(value: Any) -> str:
    return re.sub(r"\s+", " ", clean(value)).strip().lower()


def find_excel_file() -> Path:
    files = sorted(
        p for p in INPUT_DIR.iterdir()
        if p.is_file() and p.suffix.lower() in {".xlsx", ".xls"}
    )

    if not files:
        raise RuntimeError(
            f"В каталоге {INPUT_DIR} нет XLS/XLSX файла"
        )

    if len(files) > 1:
        names = ", ".join(p.name for p in files)
        raise RuntimeError(
            f"В каталоге {INPUT_DIR} найдено несколько Excel-файлов: {names}"
        )

    return files[0]


def find_header_row(ws) -> tuple[int, dict[str, int]]:
    best_row = None
    best_score = 0
    best_columns = {}

    for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
        columns = {}

        for col_no, value in enumerate(row, start=1):
            normalized = normalize_header(value)

            if not normalized:
                continue

            for known in KNOWN_HEADERS:
                if normalize_header(known) == normalized:
                    columns[known] = col_no
                    break

        score = len(columns)

        if score > best_score:
            best_score = score
            best_row = row_no
            best_columns = columns

    if best_row is None or best_score < 8:
        raise RuntimeError(
            f"Не удалось надёжно определить строку заголовков. "
            f"Максимальное совпадение: {best_score} колонок."
        )

    return best_row, best_columns


def get_value(row: tuple[Any, ...], columns: dict[str, int], name: str) -> str:
    col_no = columns.get(name)

    if not col_no:
        return ""

    if col_no > len(row):
        return ""

    return clean(row[col_no - 1])


def parse_coordinates(value: str) -> tuple[float | None, float | None]:
    if not value:
        return None, None

    text = value.replace(";", ",").strip()

    parts = [x.strip() for x in text.split(",")]

    if len(parts) != 2:
        return None, None

    try:
        lat = float(parts[0].replace(" ", ""))
        lng = float(parts[1].replace(" ", ""))
    except ValueError:
        return None, None

    if not (-90 <= lat <= 90 and -180 <= lng <= 180):
        return None, None

    return lat, lng


def is_empty_surface(item: dict[str, Any]) -> bool:
    fields = (
        "city",
        "address",
        "name",
        "side",
        "number",
        "type",
        "size",
        "price",
        "photo",
        "map",
    )

    return not any(item[field] for field in fields)


def build_item(
    row: tuple[Any, ...],
    columns: dict[str, int],
) -> dict[str, Any]:

    city = (
        get_value(row, columns, "Город")
        or get_value(row, columns, "Город 1")
        or get_value(row, columns, "город 2")
    )

    address = (
        get_value(row, columns, "Адрес")
        or get_value(row, columns, "Адресс 2")
    )

    name = get_value(row, columns, "Наименование")
    side = get_value(row, columns, "Сторона")
    item_type = get_value(row, columns, "Формат конструкции")
    size = get_value(row, columns, "Размер")
    price = get_value(row, columns, "Цена за месяц")
    number = get_value(row, columns, "№ конструкции")
    photo = get_value(row, columns, "Ссылка на фото")
    map_url = get_value(row, columns, "Карта")

    coordinates = get_value(row, columns, "Координаты")
    lat, lng = parse_coordinates(coordinates)

    months = {
        month: get_value(row, columns, month)
        for month in MONTHS
    }

    # Пока district определяется из города.
    # Реальную логику заполнения district перенесём на этап
    # JSON -> PostgreSQL, где можно использовать существующую
    # адресную логику AlbionAI.
    district = city

    statuses = {
        month: normalize_status(months[month])
        for month in MONTHS
    }

    return {
        "address": address,
        "city": city,
        "district": district,
        "lat": lat,
        "lng": lng,
        "map": map_url,
        "months": months,
        "name": name,
        "number": number,
        "photo": photo,
        "price": price,
        "side": side,
        "size": size,
        "statuses": statuses,
        "type": item_type,
    }


def normalize_status(value: str) -> str:
    text = clean(value)

    if not text:
        return "Свободно"

    lowered = text.lower()

    if "продано" in lowered:
        return "Занято"

    if "резерв" in lowered:
        return "В резерве"

    if "утверждени" in lowered:
        return "На утверждении"

    if re.fullmatch(r"\\d+(?:[.,]\\d+)?%", text):
        return "Доступность"

    return "Доступность"


def convert() -> Path:
    source = find_excel_file()

    print(f"Источник: {source}")
    print("Формат: XLSX")

    wb = load_workbook(
        source,
        read_only=True,
        data_only=True,
    )

    try:
        if not wb.worksheets:
            raise RuntimeError("В книге нет листов")

        ws = wb.worksheets[0]

        print(f"Лист: {ws.title}")
        print(f"Размер: {ws.max_row} строк × {ws.max_column} колонок")

        header_row, columns = find_header_row(ws)

        print(f"Строка заголовков: {header_row}")
        print("Найденные колонки:")

        for name, col_no in sorted(columns.items(), key=lambda x: x[1]):
            print(f"  {col_no}: {name}")

        records = []

        for row in ws.iter_rows(
            min_row=header_row + 1,
            values_only=True,
        ):
            item = build_item(row, columns)

            if is_empty_surface(item):
                continue

            records.append(item)

    finally:
        wb.close()

    output = TMP_DIR / "import.json"

    payload = {
        "source": {
            "filename": source.name,
            "format": source.suffix.lower().lstrip("."),
            "header_row": header_row,
            "year": 2026,
        },
        "count": len(records),
        "months": MONTHS,
        "surfaces": records,
    }

    output.write_text(
        json.dumps(
            payload,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    print()
    print("===== RESULT =====")
    print(f"Записей: {len(records)}")
    print(f"JSON: {output}")

    return output


if __name__ == "__main__":
    convert()
