from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .json_to_postgres import import_year

PROJECT_DIR = Path(__file__).resolve().parents[2]
INPUT_DIR = PROJECT_DIR / "input"
TMP_DIR = PROJECT_DIR / "tmp"

MONTHS = ["Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]

KNOWN_HEADERS = {
    "Город",
    "Город 1",
    "город 2",
    "Район",
    "Наименование",
    "Адрес",
    "Адресс 2",
    "Сторона",
    "Формат конструкции",
    "Размер",
    "Цена за месяц",
    "Координаты",
    "Свет",
    "Материал",
    "№ конструкции",
    "Ссылка на фото",
    "Фото",
    "Карта",
    *MONTHS,
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_header(value: Any) -> str:
    return re.sub(r"\s+", " ", clean(value)).strip().lower()


def normalize_hash_value(value: Any) -> str:
    text = clean(value).lower()
    return re.sub(r"\s+", " ", text)


def build_surface_hash(
    *,
    city: str,
    district: str,
    address: str,
    name: str,
    side: str,
    external_number: str,
    surface_type: str,
    size: str,
    latitude: float | None,
    longitude: float | None,
) -> str:
    parts = [
        normalize_hash_value(city),
        normalize_hash_value(district),
        normalize_hash_value(address),
        normalize_hash_value(name),
        normalize_hash_value(side),
        normalize_hash_value(external_number),
        normalize_hash_value(surface_type),
        normalize_hash_value(size),
        "" if latitude is None else f"{latitude:.8f}",
        "" if longitude is None else f"{longitude:.8f}",
    ]

    canonical = "\x1f".join(parts)
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def find_excel_file() -> Path:
    files = sorted(
        p for p in INPUT_DIR.iterdir()
        if p.is_file() and p.suffix.lower() in {".xlsx", ".xls"}
    )

    if not files:
        raise RuntimeError(
            f"В каталоге {INPUT_DIR} нет XLS/XLSX файла"
        )

    if len(files) > 1:
        names = ", ".join(p.name for p in files)
        raise RuntimeError(
            f"В каталоге {INPUT_DIR} найдено несколько Excel-файлов: {names}"
        )

    return files[0]


def find_header_row(ws) -> tuple[int, dict[str, int]]:
    best_row = None
    best_score = 0
    best_columns: dict[str, int] = {}

    for row_no, row in enumerate(
        ws.iter_rows(values_only=True),
        start=1,
    ):
        columns: dict[str, int] = {}

        for col_no, value in enumerate(row, start=1):
            normalized = normalize_header(value)

            if not normalized:
                continue

            for known in KNOWN_HEADERS:
                if normalize_header(known) == normalized:
                    columns[known] = col_no
                    break

        score = len(columns)

        if score > best_score:
            best_score = score
            best_row = row_no
            best_columns = columns

    if best_row is None or best_score < 8:
        raise RuntimeError(
            "Не удалось надёжно определить строку заголовков. "
            f"Максимальное совпадение: {best_score} колонок."
        )

    return best_row, best_columns


# В выгрузке 1С ссылки на фото/карту лежат формулой:
#   =HYPERLINK("https://disk.yandex.ru/i/…","Ссылка")
# При чтении с data_only=True openpyxl вернул бы текст «Ссылка» вместо
# адреса (и медиа не скачалось бы). Поэтому читаем формулы
# (data_only=False) и достаём из них сам URL.
_HYPERLINK_RE = re.compile(r'^=HYPERLINK\(\s*"([^"]+)"', re.IGNORECASE)


def extract_hyperlink(value: Any) -> str:
    if not isinstance(value, str):
        return ""

    text = value.strip()

    if not text.upper().startswith("=HYPERLINK"):
        return ""

    match = _HYPERLINK_RE.match(text)
    return match.group(1).strip() if match else ""


def get_value(
    row: tuple[Any, ...],
    columns: dict[str, int],
    name: str,
) -> str:
    col_no = columns.get(name)

    if not col_no or col_no > len(row):
        return ""

    raw = row[col_no - 1]

    # Ячейка-ссылка → отдаём адрес, а не подпись «Ссылка».
    url = extract_hyperlink(raw)

    if url:
        return url

    return clean(raw)


def parse_coordinates(
    value: str,
) -> tuple[float | None, float | None]:
    if not value:
        return None, None

    text = value.replace(";", ",").strip()
    parts = [x.strip() for x in text.split(",")]

    if len(parts) != 2:
        return None, None

    try:
        lat = float(parts[0].replace(" ", ""))
        lng = float(parts[1].replace(" ", ""))
    except ValueError:
        return None, None

    if not (-90 <= lat <= 90 and -180 <= lng <= 180):
        return None, None

    return lat, lng


def normalize_status(value: str) -> str:
    text = clean(value)

    if not text:
        return "Свободно"

    lowered = text.lower()

    if "продано" in lowered:
        return "Занято"

    if "резерв" in lowered:
        return "В резерве"

    if "утверждени" in lowered:
        return "На утверждении"

    if re.fullmatch(r"\d+(?:[.,]\d+)?%", text):
        return "Доступность"

    return "Доступность"


def is_empty_surface(item: dict[str, Any]) -> bool:
    fields = (
        "city",
        "address",
        "name",
        "side",
        "external_number",
        "type",
        "size",
        "price",
        "photo_url",
        "map_url",
    )

    return not any(item[field] for field in fields)


def build_surface(
    row: tuple[Any, ...],
    columns: dict[str, int],
) -> dict[str, Any]:

    city = (
        get_value(row, columns, "Город")
        or get_value(row, columns, "Город 1")
        or get_value(row, columns, "город 2")
    )

    address = (
        get_value(row, columns, "Адрес")
        or get_value(row, columns, "Адресс 2")
    )

    coordinates = get_value(row, columns, "Координаты")
    lat, lng = parse_coordinates(coordinates)

    months = {
        month: get_value(row, columns, month)
        for month in MONTHS
    }

    statuses = {
        month: normalize_status(months[month])
        for month in MONTHS
    }

    district = get_value(row, columns, "Район") or city
    name = get_value(row, columns, "Наименование")
    side = get_value(row, columns, "Сторона")
    external_number = get_value(row, columns, "№ конструкции")
    surface_type = get_value(row, columns, "Формат конструкции")
    size = get_value(row, columns, "Размер")

    surface_hash = build_surface_hash(
        city=city,
        district=district,
        address=address,
        name=name,
        side=side,
        external_number=external_number,
        surface_type=surface_type,
        size=size,
        latitude=lat,
        longitude=lng,
    )

    return {
        "surface_hash": surface_hash,
        "city": city,
        "district": district,
        "address": address,
        "name": name,
        "side": side,
        "external_number": external_number,
        "type": surface_type,
        "size": size,
        "price": get_value(row, columns, "Цена за месяц"),
        "light": get_value(row, columns, "Свет"),
        "material": get_value(row, columns, "Материал"),
        "photo_url": (
            get_value(row, columns, "Ссылка на фото")
            or get_value(row, columns, "Фото")
        ),
        "map_url": get_value(row, columns, "Карта"),
        "latitude": lat,
        "longitude": lng,
        "months": months,
        "statuses": statuses,
    }


def convert() -> Path:
    source = find_excel_file()

    print(f"Источник: {source}")
    print(f"Формат: {source.suffix.lower().lstrip('.')}")

    wb = load_workbook(
        source,
        read_only=True,
        # data_only=False — чтобы читать формулы (=HYPERLINK) и доставать
        # из них адреса фото/карты (см. get_value / extract_hyperlink).
        data_only=False,
    )

    try:
        if not wb.worksheets:
            raise RuntimeError("В книге нет листов")

        ws = wb.worksheets[0]

        print(f"Лист: {ws.title}")
        print(
            f"Размер: {ws.max_row} строк × "
            f"{ws.max_column} колонок"
        )

        header_row, columns = find_header_row(ws)

        print(f"Строка заголовков: {header_row}")
        print("Найденные колонки:")

        for name, col_no in sorted(
            columns.items(),
            key=lambda x: x[1],
        ):
            print(f"  {col_no}: {name}")

        records = []

        for row in ws.iter_rows(
            min_row=header_row + 1,
            values_only=True,
        ):
            item = build_surface(row, columns)

            if is_empty_surface(item):
                continue

            records.append(item)

    finally:
        wb.close()

    TMP_DIR.mkdir(parents=True, exist_ok=True)

    output = TMP_DIR / "constructions_import.json"

    payload = {
        "source": {
            "filename": source.name,
            "format": source.suffix.lower().lstrip("."),
            "header_row": header_row,
            "year": import_year(),
        },
        "count": len(records),
        "months": MONTHS,
        "surfaces": records,
    }

    output.write_text(
        json.dumps(
            payload,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    print()
    print("===== CONSTRUCTIONS RESULT =====")
    print(f"Записей: {len(records)}")
    print(f"JSON: {output}")

    return output


if __name__ == "__main__":
    convert()
